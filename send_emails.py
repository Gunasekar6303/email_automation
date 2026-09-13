import base64
import os
import time
from email.message import EmailMessage

import openpyxl
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build


# ============================================================
# CONFIGURATION
# ============================================================

EXCEL_FILE = r"D:\Projects\email_automation\Sheet1.xlsx"

RESUME_FILE = r"D:\Projects\email_automation\Gunasekar R.pdf"

CREDENTIALS_FILE = r"D:\Projects\email_automation\credentials.json"

TOKEN_FILE = r"D:\Projects\email_automation\token.json"

MY_EMAIL = "gunasekar1652@gmail.com"

MAX_EMAILS_PER_RUN = None

SCOPES = [
    "https://www.googleapis.com/auth/gmail.send"
]


# ============================================================
# GMAIL AUTHENTICATION
# ============================================================

def get_gmail_service():
    """Authenticate with Gmail API and return the Gmail service."""

    credentials = None

    if os.path.exists(TOKEN_FILE):
        credentials = Credentials.from_authorized_user_file(
            TOKEN_FILE,
            SCOPES
        )

    if not credentials or not credentials.valid:

        if credentials and credentials.expired and credentials.refresh_token:
            credentials.refresh(Request())

        else:
            flow = InstalledAppFlow.from_client_secrets_file(
                CREDENTIALS_FILE,
                SCOPES
            )

            credentials = flow.run_local_server(
                port=0
            )

        with open(TOKEN_FILE, "w", encoding="utf-8") as token_file:
            token_file.write(credentials.to_json())

    return build(
        "gmail",
        "v1",
        credentials=credentials
    )


# ============================================================
# CREATE EMAIL
# ============================================================

def create_email(
    recipient,
    hr_name,
    resume_path
):
    """Create email with resume attachment."""

    if hr_name:
        greeting = f"Hi {hr_name},"
    else:
        greeting = "Hi Hiring Manager,"

    subject = "AI Engineer – GenAI/RAG | 2 Years Experience"

    body = f"""{greeting}

I’m Gunasekar R, an AI/ML Engineer with 2 years of hands-on
experience building AI/ML and Generative AI applications.

I’m currently working on solutions involving LLMs, RAG, Agentic AI,
NLP, Computer Vision, and Machine Learning. I also have prior
Computer Vision project experience with the IIT Madras Incubation Cell.

My experience includes building and deploying AI applications using
Python, LangChain, LangGraph, LLMs, RAG, FastAPI, Docker, Kubernetes,
AWS, and vector databases, along with ML/CV solutions for real-world
use cases.

I’m currently exploring AI/ML, GenAI, and Agentic AI Engineer
opportunities where I can contribute to production AI systems and
continue growing technically.

I’ve attached my resume for your consideration. If there is a
suitable opening in your team, I’d be grateful for the opportunity
to discuss it.

Best Regards,

Gunasekar R
AI/ML Engineer

+91-9345515578
gunasekar1652@gmail.com

LinkedIn:
https://www.linkedin.com/in/gunasekar1652/

GitHub:
https://github.com/Gunasekar6303
"""

    message = EmailMessage()

    message["To"] = recipient
    message["Subject"] = subject

    message.set_content(body)

    with open(resume_path, "rb") as resume_file:
        resume_data = resume_file.read()

    message.add_attachment(
        resume_data,
        maintype="application",
        subtype="pdf",
        filename=os.path.basename(resume_path)
    )

    encoded_message = base64.urlsafe_b64encode(
        message.as_bytes()
    ).decode()

    return {
        "raw": encoded_message
    }


# ============================================================
# SEND EMAIL
# ============================================================

def send_email(
    gmail_service,
    recipient,
    hr_name,
    resume_path
):
    """Send one email through Gmail API."""

    message = create_email(
        recipient,
        hr_name,
        resume_path
    )

    gmail_service.users().messages().send(
        userId="me",
        body=message
    ).execute()


# ============================================================
# MAIN
# ============================================================

def main():

    # --------------------------------------------------------
    # CHECK FILES
    # --------------------------------------------------------

    if not os.path.exists(EXCEL_FILE):
        print(f"ERROR: Excel file not found:")
        print(EXCEL_FILE)
        return

    if not os.path.exists(RESUME_FILE):
        print(f"ERROR: Resume file not found:")
        print(RESUME_FILE)
        return

    if not os.path.exists(CREDENTIALS_FILE):
        print(f"ERROR: credentials.json not found:")
        print(CREDENTIALS_FILE)
        return

    # --------------------------------------------------------
    # AUTHENTICATE
    # --------------------------------------------------------

    print("Connecting to Gmail...")

    gmail_service = get_gmail_service()

    print("Gmail authentication successful.")
    print()

    # --------------------------------------------------------
    # LOAD EXCEL
    # --------------------------------------------------------

    workbook = openpyxl.load_workbook(
        EXCEL_FILE
    )

    sheet = workbook.active

    emails_sent = 0

    # --------------------------------------------------------
    # PROCESS HR LIST
    # --------------------------------------------------------

    for row in range(2, sheet.max_row + 1):

        email = sheet.cell(
            row=row,
            column=1
        ).value

        hr_name = sheet.cell(
            row=row,
            column=2
        ).value

        status = sheet.cell(
            row=row,
            column=3
        ).value

        # ----------------------------------------------------
        # SKIP EMPTY ROW
        # ----------------------------------------------------

        if not email:
            continue

        email = str(email).strip()

        # ----------------------------------------------------
        # NEVER SEND TO YOURSELF
        # ----------------------------------------------------

        if email.lower() == MY_EMAIL.lower():

            print(f"Skipping your own email: {email}")

            continue

        # ----------------------------------------------------
        # SKIP ALREADY SENT
        # ----------------------------------------------------

        if str(status).strip().upper() == "SENT":

            print(f"Already sent, skipping: {email}")

            continue

        # ----------------------------------------------------
        # BATCH LIMIT
        # ----------------------------------------------------

        if MAX_EMAILS_PER_RUN is not None and emails_sent >= MAX_EMAILS_PER_RUN:

            break

        print(f"Sending to: {email}")

        try:

            send_email(
                gmail_service,
                email,
                hr_name,
                RESUME_FILE
            )

            # ------------------------------------------------
            # SUCCESS
            # ------------------------------------------------

            sheet.cell(
                row=row,
                column=3
            ).value = "SENT"

            workbook.save(EXCEL_FILE)

            emails_sent += 1

            print(f"SUCCESS: {email} -> SENT")
            print()

            # Small delay
            time.sleep(3)

        except Exception as error:

            # ------------------------------------------------
            # FAILURE
            # ------------------------------------------------

            sheet.cell(
                row=row,
                column=3
            ).value = "FAILED"

            workbook.save(EXCEL_FILE)

            print(f"FAILED: {email}")
            print(f"Error: {error}")
            print()

    # --------------------------------------------------------
    # FINAL SUMMARY
    # --------------------------------------------------------

    print("==============================")
    print(f"Emails sent: {emails_sent}")
    print("==============================")


# ============================================================
# START PROGRAM
# ============================================================

if __name__ == "__main__":
    main()